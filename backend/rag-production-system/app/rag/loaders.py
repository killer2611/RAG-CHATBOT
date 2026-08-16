from __future__ import annotations

from pathlib import Path

from langchain_core.documents import Document
from pypdf import PdfReader
from docx import Document as DocxDocument

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx"}


def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix or '<none>'}")

    if suffix == ".pdf":
        reader = PdfReader(str(path))
        docs: list[Document] = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                docs.append(
                    Document(
                        page_content=text,
                        metadata={"source_name": path.name, "page": page_number},
                    )
                )
        return docs

    if suffix == ".docx":
        docx = DocxDocument(str(path))
        paragraphs = [p.text.strip() for p in docx.paragraphs if p.text.strip()]
        text = "\n".join(paragraphs)
        return [Document(page_content=text, metadata={"source_name": path.name})] if text else []

    text = path.read_text(encoding="utf-8", errors="replace")
    return [Document(page_content=text, metadata={"source_name": path.name})] if text.strip() else []
